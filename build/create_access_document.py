from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
source_accounts = (ROOT / 'docs/demo-accounts.md').read_text(encoding='utf-8')
source_status = (ROOT / 'docs/implementation-status.md').read_text(encoding='utf-8')
accounts = []
for line in source_accounts.splitlines():
    if line.startswith('| ') and '@caresys.test' in line:
        accounts.append([v.strip() for v in line.strip('|').split('|')])
assert len(accounts) == 12
roles = ['Platform administrator', 'Agency owner', 'Recruiter', 'Scheduler', 'Supervisor', 'Finance officer',
         'Caregiver', 'Independent candidate', 'Former employee', 'Identity review candidate', 'Agency owner', 'Caregiver']
headings = ['Identity email and signing', 'Subscriptions and payments', 'References documents and privacy',
            'Scheduling and workforce operations', 'Payroll processing', 'Billing and accounting',
            'Production deployment and assurance', 'Policy decisions and restricted features']
section = source_status.split('## Remaining work and external dependencies', 1)[1].split('\n## ', 1)[0]
items = []
for line in section.splitlines():
    if line.startswith('- '):
        items.append(line[2:].strip())
    elif line.strip() and items:
        items[-1] += ' ' + line.strip()
assert len(items) == 8

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.65), Inches(.65)
s.left_margin, s.right_margin = Inches(.65), Inches(.65)
s.footer_distance = Inches(.3)
styles = doc.styles
for style in styles:
    for borders in list(style.element.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.07
styles['Title'].font.size = Pt(24)
styles['Title'].paragraph_format.space_after = Pt(7)
styles['Heading 1'].font.size = Pt(18)
styles['Heading 1'].paragraph_format.space_after = Pt(8)
styles['Heading 2'].font.size = Pt(11.5)
styles['Heading 2'].paragraph_format.space_before = Pt(8)
styles['Heading 2'].paragraph_format.space_after = Pt(3)
doc.core_properties.title = 'CareSys Demo Accounts and Remaining Work'
doc.core_properties.subject = 'Local demonstration access and outstanding implementation requirements'
doc.core_properties.author = 'CareSys'

doc.add_paragraph('CareSys Demo Accounts\nand Remaining Work', 'Title')
p = doc.add_paragraph('23 October 2026  |  Local development access')
p.paragraph_format.space_after = Pt(9)
doc.add_paragraph('Use these accounts to review the local CareSys application. The remaining work on the next page identifies requirements that are still unfinished or depend on external services and policy decisions.')
p = doc.add_paragraph()
p.add_run('Login: ').bold = True
p.add_run('http://localhost:8080/login')
p.add_run('    Shared password: ').bold = True
p.add_run('CareSysDemo!2026').bold = True

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [2.00, 1.35, 3.85]
for c, w in zip(table.columns, widths): c.width = Inches(w)
for cell, label in zip(table.rows[0].cells, ['Email account', 'Role', 'Workspace and access']): cell.text = label
repeat = OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(repeat)
for idx, (email, access) in enumerate(accounts):
    cells = table.add_row().cells
    for c, value in zip(cells, [email, roles[idx], access]): c.text = value
for ri, row in enumerate(table.rows):
    no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    for ci, cell in enumerate(row.cells):
        cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for side, amount in [('top', 65), ('bottom', 65), ('left', 85), ('right', 85)]:
            el = OxmlElement('w:'+side); el.set(qn('w:w'), str(amount)); el.set(qn('w:type'), 'dxa'); margins.append(el)
        pr.append(margins)
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '35434A' if ri==0 else ('F3F5F6' if ri%2==0 else 'FFFFFF')); pr.append(shade)
        borders = OxmlElement('w:tcBorders')
        for side in ['top', 'left', 'bottom', 'right']:
            edge = OxmlElement('w:'+side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
        pr.append(borders)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            for r in p.runs:
                r.font.size = Pt(9)
                if ri==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
p=doc.add_paragraph('All people and records are fictional. Demo verification and subscription states do not represent actual checks or payments. Internal service accounts use separate random passwords. The shared demo password is for local development only.')
p.paragraph_format.space_before = Pt(8)
for r in p.runs: r.font.size=Pt(9)

doc.add_page_break()
doc.add_paragraph('Remaining Work Checklist', 'Heading 1')
doc.add_paragraph('The principal local workflows are implemented. The following eight work areas remain open; no external payment, certified signature, completed screening, or production readiness is implied.')
for idx, (heading, body) in enumerate(zip(headings, items), 1):
    doc.add_paragraph(f'{idx}  {heading}', 'Heading 2')
    p=doc.add_paragraph(body)
    p.paragraph_format.keep_together=True
doc.add_paragraph('Completion evidence', 'Heading 2')
doc.add_paragraph('Close each work area only after the applicable provider or policy is approved, the implementation is tested, and the result is verified. Expanded concurrency tests, formal accessibility review, and a measured restore exercise remain necessary.')

footer=s.footer.paragraphs[0]
footer.alignment=2
r=footer.add_run('CareSys  |  '); r.font.size=Pt(8);r.font.color.rgb=RGBColor(90,90,90)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
out=ROOT/'CareSys_Demo_Accounts_and_Remaining_Work.docx'
doc.save(out)
print(out)
